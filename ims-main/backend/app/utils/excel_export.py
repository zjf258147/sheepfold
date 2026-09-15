"""通用 Excel 导出工具。"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def build_simple_xlsx(
    headers: list[str],
    rows: list[list],
    sheet_title: str = "Sheet1",
    col_widths: list[int] | None = None,
) -> bytes:
    """将表头+行数据生成为 Excel（xlsx）字节流。

    Args:
        headers: 表头列名列表
        rows: 数据行列表，每行为 list
        sheet_title: Sheet 名称
        col_widths: 列宽列表（可选）
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="409EFF")
    center = Alignment(horizontal="center", vertical="center")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center

    for row in rows:
        ws.append(row)

    if col_widths:
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()