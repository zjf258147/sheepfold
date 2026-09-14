"""Excel 导入导出专项测试（P1）

覆盖：
  1. 导出内容与数据库一致
  2. 长数字文本不被转科学计数法
  3. 物料编码匹配失败友好提示
  4. 多行名称解析

运行方式：
  cd backend && python -m pytest tests/test_excel.py -v --tb=short
"""

import io
import pytest
from datetime import date, datetime
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from app.utils.excel_export import build_simple_xlsx
from app.utils.excel_import import parse_import_xlsx, build_template_xlsx
from app.utils.item_export import (
    build_items_xlsx,
    build_ledger_xlsx,
    build_ledger_summary_xlsx,
    stock_status_label,
    stock_condition_label,
    operation_status_label,
)


# ============================================================
# 1. 导出内容与数据库一致
# ============================================================

class TestExportContentIntegrity:
    """验证导出 Excel 的内容与数据源一致。"""

    def test_build_simple_xlsx_headers_and_rows(self):
        """导出的 Excel 表头和数据行与输入一致。"""
        headers = ["编号", "名称", "数量", "日期"]
        rows = [
            ["A001", "物料A", "10", "2024-01-01"],
            ["A002", "物料B", "20", "2024-01-02"],
        ]
        data = build_simple_xlsx(headers, rows, sheet_title="测试")

        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.title == "测试"

        assert ws.cell(1, 1).value == "编号"
        assert ws.cell(1, 2).value == "名称"
        assert ws.cell(1, 3).value == "数量"
        assert ws.cell(1, 4).value == "日期"

        assert ws.cell(2, 1).value == "A001"
        assert ws.cell(2, 2).value == "物料A"
        assert ws.cell(2, 3).value == "10"
        assert ws.cell(2, 4).value == "2024-01-01"

        assert ws.cell(3, 1).value == "A002"
        assert ws.cell(3, 4).value == "2024-01-02"

    def test_build_simple_xlsx_empty_data(self):
        """空数据导出时仅包含表头。"""
        data = build_simple_xlsx(["A", "B"], [], sheet_title="空表")
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.max_row == 1
        assert ws.cell(1, 1).value == "A"
        assert ws.cell(1, 2).value == "B"

    def test_build_simple_xlsx_header_style(self):
        """表头样式：蓝色背景、白色粗体。"""
        data = build_simple_xlsx(["列1"], [["值1"]])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        cell = ws.cell(1, 1)
        assert cell.font.bold is True
        assert cell.font.color.rgb in ("00FFFFFF", "FFFFFF")
        assert cell.fill.fgColor.rgb in ("00409EFF", "409EFF")

    def test_build_items_xlsx_content_match(self):
        """单品明细导出内容与输入数据一致。"""
        items = [
            {
                "item_sn": "SN20240001",
                "sku_name": "测试产品A",
                "unit_price": 1234.56,
                "stock_status": "IN_STOCK",
                "stock_condition": "NEW",
                "last_order_no": "JIN-001",
                "partner_group_name": "",
                "partner_name": "",
                "customer_name": "",
                "operation_status": "COMPLETED",
                "remark": "测试备注",
            },
            {
                "item_sn": "SN20240002",
                "sku_name": "测试产品B",
                "unit_price": 99.00,
                "stock_status": "SOLD",
                "stock_condition": "SOLD",
                "last_order_no": "JOUT-001",
                "partner_group_name": "客户分组",
                "partner_name": "客户A",
                "customer_name": "客户A",
                "operation_status": "COMPLETED",
                "remark": "",
            },
        ]
        data = build_items_xlsx(items, sheet_title="库存明细")
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.title == "库存明细"

        assert ws.cell(2, 1).value == "SN20240001"
        assert ws.cell(2, 2).value == "测试产品A"
        assert ws.cell(2, 3).value == 1234.56
        assert ws.cell(2, 4).value == "在库"
        assert ws.cell(2, 5).value == "正常"
        assert ws.cell(2, 6).value == "JIN-001"
        assert ws.cell(2, 10).value == "已完成"
        assert ws.cell(2, 11).value == "测试备注"

        assert ws.cell(3, 1).value == "SN20240002"
        assert ws.cell(3, 4).value == "出库"
        assert ws.cell(3, 5).value == "售出-线上"
        assert ws.cell(3, 7).value == "客户分组"
        assert ws.cell(3, 8).value == "客户A"
        assert ws.cell(3, 9).value == "客户A"

    def test_build_items_xlsx_freeze_panes(self):
        """冻结窗格设置在 A2。"""
        data = build_items_xlsx([{"item_sn": "SN001", "sku_name": "X"}])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.freeze_panes == "A2"

    def test_item_lookup_functions(self):
        """库存状态/属性/操作状态转换函数正确。"""
        assert stock_status_label("IN_STOCK") == "在库"
        assert stock_status_label("OUT_STOCK") == "出库"
        assert stock_status_label(None) == "出库"

        assert stock_condition_label("IN_STOCK", "NEW") == "正常"
        assert stock_condition_label("IN_STOCK", "RETURNED_FROM_SALE") == "线上售出退回"
        assert stock_condition_label("SOLD", None) == "售出-线上"
        assert stock_condition_label("BORROWED", None) == "已借用"

        assert operation_status_label("INITIATED") == "已发起"
        assert operation_status_label("COMPLETED") == "已完成"
        assert operation_status_label("CANCELLED") == "已取消"
        assert operation_status_label(None) == ""


class TestLedgerExportContent:
    """库存日流水/快照汇总导出内容验证。"""

    def test_build_ledger_xlsx_content(self):
        """日流水导出内容正确。"""
        rows = [
            {
                "snapshot_date": "2024-01-15",
                "sku_name": "产品A",
                "opening_in_stock_qty": 100,
                "inbound_by_type": {"INBOUND": 10, "RETURN": 5},
                "outbound_by_type": {"SOLD": 8, "SCRAPPED": 2},
                "closing_in_stock_qty": 105,
                "closing_asset_amount": 52500.00,
            },
        ]
        data = build_ledger_xlsx(rows)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.title == "库存日流水"
        assert ws.cell(2, 1).value == "2024-01-15"
        assert ws.cell(2, 2).value == "产品A"
        assert ws.cell(2, 3).value == 100
        assert ws.cell(2, 11).value == 105
        assert ws.cell(2, 12).value == 52500.00

    def test_build_ledger_summary_xlsx_content(self):
        """快照汇总导出内容正确。"""
        rows = [
            {
                "sku_name": "产品A",
                "opening_in_stock_qty": 50,
                "inbound_by_type": {"INBOUND": 20},
                "outbound_by_type": {"SOLD": 15},
                "expected_closing_qty": 55,
                "closing_in_stock_qty": 55,
                "closing_asset_amount": 27500.00,
                "balanced": True,
            },
            {
                "sku_name": "产品B",
                "opening_in_stock_qty": 30,
                "inbound_by_type": {},
                "outbound_by_type": {"SCRAPPED": 5},
                "expected_closing_qty": 25,
                "closing_in_stock_qty": 23,
                "closing_asset_amount": 0,
                "balanced": False,
                "diff_qty": -2,
            },
        ]
        data = build_ledger_summary_xlsx(rows, "2024-01-01", "2024-01-31")
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.title == "库存快照汇总"

        assert ws.cell(2, 1).value == "产品A"
        assert ws.cell(2, 2).value == 50
        assert ws.cell(2, 11).value == 55
        assert ws.cell(2, 13).value == "账实一致"

        assert ws.cell(3, 1).value == "产品B"
        assert ws.cell(3, 13).value == "差异 -2 件"

    def test_ledger_summary_unknown_opening(self):
        """期初未知时显示'期初未知'。"""
        rows = [
            {
                "sku_name": "新产品",
                "opening_in_stock_qty": None,
                "inbound_by_type": {},
                "outbound_by_type": {},
                "expected_closing_qty": None,
                "closing_in_stock_qty": 10,
                "closing_asset_amount": 5000.00,
                "balanced": False,
            },
        ]
        data = build_ledger_summary_xlsx(rows, "2024-01-01", "2024-01-31")
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.cell(2, 13).value == "期初未知"


# ============================================================
# 2. 长数字文本不被转科学计数法
# ============================================================

class TestLongNumberText:
    """长数字文本（条码、SN号等）在 Excel 导出/导入时不被转换为科学计数法。"""

    def test_barcode_not_scientific_notation(self):
        """长数字条码导出后保持原始文本格式。"""
        barcode = "697123456789012345678"
        data = build_simple_xlsx(["条码"], [[barcode]])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.cell(2, 1).value == barcode

    def test_sn_not_scientific_notation(self):
        """SN号（纯数字）导出后保持原始文本格式。"""
        sn = "2024011500000001"
        data = build_simple_xlsx(["SN号"], [[sn]])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.cell(2, 1).value == sn

    def test_item_sn_in_export_not_scientific(self):
        """单品明细中的 SN 号导出后保持原始格式。"""
        items = [
            {
                "item_sn": "20240115999999999999",
                "sku_name": "测试",
                "unit_price": None,
                "stock_status": "IN_STOCK",
                "stock_condition": "NEW",
                "last_order_no": "",
                "partner_group_name": "",
                "partner_name": "",
                "customer_name": "",
                "operation_status": None,
                "remark": "",
            },
        ]
        data = build_items_xlsx(items)
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.cell(2, 1).value == "20240115999999999999"

    def test_parse_import_strips_whitespace(self):
        """导入时自动去除前后空格。"""
        data = build_simple_xlsx(["编码", "名称"], [["  SKU001  ", "  测试物料  "]])
        rows = parse_import_xlsx(data)
        assert rows[0] == ["SKU001", "测试物料"]

    def test_import_empty_rows_skipped(self):
        """导入时跳过空行。"""
        data = build_simple_xlsx(["编码", "名称"], [["SKU001", "物料A"], ["", ""], ["SKU002", "物料B"]])
        rows = parse_import_xlsx(data)
        assert len(rows) == 2
        assert rows[0] == ["SKU001", "物料A"]
        assert rows[1] == ["SKU002", "物料B"]

    def test_unicode_preserved_in_export_import(self):
        """中文、特殊字符在导出→导入循环中保持一致。"""
        test_name = "测试物料-特殊字符©®™"
        data = build_simple_xlsx(["名称"], [[test_name]])
        rows = parse_import_xlsx(data)
        assert rows[0] == [test_name]


# ============================================================
# 3. 物料编码匹配失败友好提示
# ============================================================

class TestMaterialCodeMatching:
    """物料编码（SKU Code）在导入时匹配失败给出友好提示。"""

    def test_bom_import_missing_sku_code(self):
        """BOM导入时物料编码不存在，给出明确提示。"""
        from app.service.bom_service import import_bom_xlsx

        mock_db = MagicMock()
        mock_db.query.return_value.filter.return_value.first.return_value = None

        data = build_simple_xlsx(
            ["BOM名称", "版本号", "物料编码", "物料名称", "计划数量", "备注"],
            [["BOM1", "V1.0", "SKU999", "不存在的物料", "10", ""]],
        )
        result = import_bom_xlsx(mock_db, data, "admin")
        assert result["success"] == 0
        assert len(result["errors"]) == 1
        assert "SKU999" in result["errors"][0]
        assert "不存在" in result["errors"][0]

    def test_bom_import_valid_sku_creates_record(self):
        """BOM导入时物料编码存在，正常创建记录。"""
        from app.service.bom_service import import_bom_xlsx
        from app.models.product import ProductSku

        mock_sku = MagicMock(spec=ProductSku)
        mock_sku.id = 1
        mock_sku.sku_code = "SKU001"
        mock_sku.name = "测试物料"

        mock_db = MagicMock()
        mock_db.query.return_value.filter.return_value.first.return_value = mock_sku

        data = build_simple_xlsx(
            ["BOM名称", "版本号", "物料编码", "物料名称", "计划数量", "备注"],
            [["BOM1", "V1.0", "SKU001", "测试物料", "100", "备注"]],
        )
        result = import_bom_xlsx(mock_db, data, "admin")
        assert result["success"] == 1
        assert len(result["errors"]) == 0
        mock_db.add.assert_called_once()
        mock_db.commit.assert_called_once()

    def test_bom_import_empty_name(self):
        """BOM导入时名称为空给出提示。"""
        from app.service.bom_service import import_bom_xlsx

        mock_db = MagicMock()
        data = build_simple_xlsx(
            ["BOM名称", "版本号", "物料编码", "物料名称", "计划数量", "备注"],
            [["", "V1.0", "SKU001", "物料", "10", ""]],
        )
        result = import_bom_xlsx(mock_db, data, "admin")
        assert result["success"] == 0
        assert len(result["errors"]) == 1
        assert "名称为空" in result["errors"][0]

    def test_bom_import_insufficient_columns(self):
        """BOM导入时关键字段为空给出提示。"""
        from app.service.bom_service import import_bom_xlsx

        mock_db = MagicMock()
        mock_db.query.return_value.filter.return_value.first.return_value = None

        data = build_simple_xlsx(
            ["BOM名称", "版本号", "物料编码", "物料名称", "计划数量", "备注"],
            [["BOM1", "V1.0", "", "", "", ""]],
        )
        result = import_bom_xlsx(mock_db, data, "admin")
        assert result["success"] == 0
        assert len(result["errors"]) == 1
        assert "不存在" in result["errors"][0]

    def test_sku_import_duplicate_barcode(self):
        """SKU导入时条码重复给出友好提示。"""
        from app.service.product_service import import_sku_xlsx

        mock_db = MagicMock()
        mock_db.query.return_value.filter.return_value.first.return_value = MagicMock()

        data = build_simple_xlsx(
            ["物料名称", "分类名称", "物料编码", "规格型号", "条码", "单位", "SN模式", "物料类型", "备注"],
            [["物料A", "成品", "SKU001", "规格", "BAR001", "个", "BOTH", "成品", ""]],
        )
        result = import_sku_xlsx(mock_db, data)
        assert result["success"] == 0
        assert len(result["errors"]) == 1
        assert "BAR001" in result["errors"][0]
        assert "已存在" in result["errors"][0]

    def test_partner_import_empty_name(self):
        """往来单位导入时名称为空给出提示。"""
        from app.service.partner_service import import_partner_xlsx

        mock_db = MagicMock()
        data = build_simple_xlsx(
            ["单位名称", "分组名称", "类型", "备注"],
            [["", "分组", "供应商", ""]],
        )
        result = import_partner_xlsx(mock_db, data)
        assert result["success"] == 0
        assert len(result["errors"]) == 1
        assert "名称为空" in result["errors"][0]


# ============================================================
# 4. 多行名称解析
# ============================================================

class TestMultiLineName:
    """多行名称（含换行符、特殊字符、长名称）解析。"""

    def test_name_with_newline(self):
        """名称含换行符时正确解析。"""
        name = "第一行\n第二行"
        data = build_simple_xlsx(["名称"], [[name]])
        rows = parse_import_xlsx(data)
        assert rows[0] == [name]

    def test_name_with_spaces(self):
        """名称含前后空格时自动去除。"""
        name = "  测试物料名称  "
        data = build_simple_xlsx(["名称"], [[name]])
        rows = parse_import_xlsx(data)
        assert rows[0] == ["测试物料名称"]

    def test_name_with_special_chars(self):
        """名称含特殊字符（括号、斜杠）正确解析。"""
        name = "物料(成品)/型号[A/B]"
        data = build_simple_xlsx(["名称"], [[name]])
        rows = parse_import_xlsx(data)
        assert rows[0] == [name]

    def test_long_name(self):
        """长名称（>100字符）正确解析。"""
        name = "A" * 200
        data = build_simple_xlsx(["名称"], [[name]])
        rows = parse_import_xlsx(data)
        assert rows[0] == [name]

    def test_chinese_punctuation_in_name(self):
        """中文标点符号（，。、；：）正确解析。"""
        name = "测试物料，型号A。规格B、备注C；"
        data = build_simple_xlsx(["名称"], [[name]])
        rows = parse_import_xlsx(data)
        assert rows[0] == [name]

    def test_empty_cell_treated_as_empty_string(self):
        """空单元格解析为空字符串。"""
        data = build_simple_xlsx(["A", "B", "C"], [["值1", None, "值3"]])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.cell(2, 2).value is None
        rows = parse_import_xlsx(data)
        assert rows[0] == ["值1", "", "值3"]

    def test_number_as_string(self):
        """数字类型的单元格转为字符串。"""
        data = build_simple_xlsx(["数量"], [[123]])
        rows = parse_import_xlsx(data)
        assert rows[0] == ["123"]

    def test_template_contains_example_row(self):
        """导入模板包含示例行。"""
        data = build_template_xlsx(
            ["名称", "编码"],
            example_row=["示例", "EX001"],
            sheet_title="导入模板",
        )
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.title == "导入模板"
        assert ws.cell(1, 1).value == "名称"
        assert ws.cell(1, 2).value == "编码"
        assert ws.cell(2, 1).value == "示例"
        assert ws.cell(2, 2).value == "EX001"

    def test_template_no_example_row(self):
        """无示例行时仅包含表头。"""
        data = build_template_xlsx(["名称", "编码"])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.max_row == 1
        assert ws.cell(1, 1).value == "名称"

    def test_template_header_style(self):
        """模板表头样式：绿色背景、白色粗体。"""
        data = build_template_xlsx(["列1"])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        cell = ws.cell(1, 1)
        assert cell.font.bold is True
        assert cell.font.color.rgb in ("00FFFFFF", "FFFFFF")
        assert cell.fill.fgColor.rgb in ("0067C23A", "67C23A")


# ============================================================
# 5. 边界与异常
# ============================================================

class TestExcelEdgeCases:
    """边界值与异常场景。"""

    def test_single_row_export(self):
        """单行数据导出。"""
        data = build_simple_xlsx(["A"], [["1"]])
        wb = load_workbook(io.BytesIO(data))
        assert wb.active.max_row == 2

    def test_large_dataset_export(self):
        """大批量数据导出（1000行）。"""
        headers = ["编号", "名称"]
        rows = [[f"SKU{i:04d}", f"物料{i}"] for i in range(1000)]
        data = build_simple_xlsx(headers, rows)
        wb = load_workbook(io.BytesIO(data))
        assert wb.active.max_row == 1001
        assert wb.active.cell(1001, 1).value == "SKU0999"

    def test_zero_and_negative_numbers(self):
        """零值和负数正常导出。"""
        data = build_simple_xlsx(["数量", "金额"], [[0, -100.50]])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.cell(2, 1).value == 0
        assert ws.cell(2, 2).value == -100.50

    def test_float_precision(self):
        """浮点数精度正确。"""
        data = build_simple_xlsx(["金额"], [[0.1 + 0.2]])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.cell(2, 1).value == pytest.approx(0.3)

    def test_col_widths_applied(self):
        """列宽设置生效。"""
        data = build_simple_xlsx(["A", "B"], [["1", "2"]], col_widths=[20, 30])
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws.column_dimensions["A"].width == 20
        assert ws.column_dimensions["B"].width == 30

    def test_parse_skip_multiple_header_rows(self):
        """跳过多个表头行。"""
        data = build_simple_xlsx(["H1"], [["H2"], ["H3"], ["D1"], ["D2"]])
        rows = parse_import_xlsx(data, skip_header=3)
        assert len(rows) == 2
        assert rows[0] == ["D1"]
        assert rows[1] == ["D2"]


# ============================================================
# 6. 循环一致性（导出 → 导入）
# ============================================================

class TestRoundTrip:
    """导出→导入循环一致性验证。"""

    def test_simple_round_trip(self):
        """简单数据导出再导入，内容一致。"""
        original = [["SKU001", "物料A", "10"], ["SKU002", "物料B", "20"]]
        data = build_simple_xlsx(["编码", "名称", "数量"], original)
        imported = parse_import_xlsx(data)
        assert imported == original

    def test_unicode_round_trip(self):
        """Unicode数据导出再导入，内容一致。"""
        original = [["中文名称", "日本語", "한국어"]]
        data = build_simple_xlsx(["语言"], original)
        imported = parse_import_xlsx(data)
        assert imported == original

    def test_template_round_trip(self):
        """模板导出（含示例行）→ 修改数据 → 导入。"""
        tmpl = build_template_xlsx(
            ["名称", "编码"],
            example_row=["示例名称", "EX001"],
        )
        wb = load_workbook(io.BytesIO(tmpl))
        ws = wb.active
        ws["A2"] = "新产品"
        ws["B2"] = "NEW001"
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        rows = parse_import_xlsx(buf.getvalue())
        assert rows[0] == ["新产品", "NEW001"]